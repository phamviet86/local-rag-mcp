"""Synthetic workbooks only; no Drive or embedding calls during extraction."""

import datetime as dt
import json
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from local_rag import xlsx_compat as compat
from local_rag.config import Settings
from local_rag.extract import _xlsx
from local_rag.extraction_errors import ExtractionError
from local_rag.mcp import MCPServer
from local_rag.service import MultiSourceRAG
from tests.test_drive_regressions import TreeDrive, value


def rewrite(path, changes):
    with ZipFile(path) as archive:
        members = [(info, archive.read(info)) for info in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for info, data in members:
            result = changes(info.filename, data)
            if result is not None:
                archive.writestr(info, result)


def workbook(path, color="9EB3C9"):
    book = Workbook()
    sheet = book.active
    sheet.title = "Third"
    sheet["A1"] = "Synthetic header"
    sheet["A1"].fill = PatternFill("solid", fgColor="FF9EB3C9")
    for coord, val, fmt in [
        ("C2", dt.date(2026, 1, 2), "yyyy-mm-dd"),
        ("D4", dt.time(12, 34), "hh:mm"),
        ("B6", dt.datetime(2026, 1, 2, 12, 34), "yyyy-mm-dd hh:mm"),
        ("A8", 0.125, "0.0%"),
        ("F9", 42.75, '"$"#,##0.00'),
    ]:
        sheet[coord] = val
        sheet[coord].number_format = fmt
    sheet["A10"] = "=1+2"
    sheet["A11"] = "=2+3"
    book.create_sheet("First")["B2"] = "Another header"
    book.create_sheet("Second")["C3"] = 7
    book.save(path)
    book.close()

    def change(name, data):
        if name == compat.STYLES_PART:
            return data.replace(b'rgb="FF9EB3C9"', f'rgb="{color}"'.encode())
        if name == "xl/worksheets/sheet1.xml":
            return data.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>")
        return data

    rewrite(path, change)
    return path


@pytest.fixture(autouse=True)
def no_network():
    with (
        patch("socket.socket.connect", side_effect=AssertionError("network forbidden")),
        patch(
            "local_rag.embeddings.OpenAIEmbeddings.embed",
            side_effect=AssertionError("embedding forbidden"),
        ),
        patch(
            "local_rag.embeddings.LocalEmbeddings.embed",
            side_effect=AssertionError("embedding forbidden"),
        ),
        patch(
            "local_rag.drive.GoogleDriveBackend.download",
            side_effect=AssertionError("Drive forbidden"),
        ),
    ):
        yield


@pytest.mark.parametrize("color", ["9EB3C9", "FF9EB3C9", "9eb3c9", "ff9eb3c9"])
def test_equivalence_and_immutability(tmp_path, color):
    baseline = workbook(tmp_path / "valid.xlsx", color)
    malformed = workbook(tmp_path / "hash.xlsx", "#" + color)
    cache = tmp_path / "cache"
    cache.mkdir()
    before = malformed.read_bytes(), malformed.stat().st_mtime_ns
    with pytest.raises(ValueError) as error:
        load_workbook(malformed, read_only=True, data_only=True)
    assert compat.color_failure(error.value)
    with patch("local_rag.extract.repair_styles", side_effect=AssertionError("no repair")):
        expected = _xlsx(baseline, cache)
    actual = _xlsx(malformed, cache)
    assert actual.text == expected.text
    assert actual.spans == expected.spans
    assert expected.metadata == {"sheets": ["Third", "First", "Second"]}
    assert actual.metadata == {
        **expected.metadata,
        "compatibility_repairs": [
            {"kind": "xlsx_rgb_hash_prefix", "part": "xl/styles.xml", "count": 1}
        ],
    }
    assert "sheet:Third/cell:A10" in [s.locator for s in actual.spans]
    assert "sheet:Third/cell:A11" not in [s.locator for s in actual.spans]
    assert "3" in actual.text
    assert before == (malformed.read_bytes(), malformed.stat().st_mtime_ns)
    assert list(cache.iterdir()) == []
    # v0.10.0 extraction contract, independently reconstructed from openpyxl.
    book = load_workbook(baseline, read_only=True, data_only=True)
    values = [
        (str(c.value).strip(), f"sheet:{s.title}/cell:{c.coordinate}")
        for s in book
        for row in s
        for c in row
        if c.value is not None
    ]
    book.close()
    assert expected.text == "\n\n".join(v for v, _ in values)
    assert [s.locator for s in expected.spans] == [loc for _, loc in values]


@pytest.mark.parametrize("color", ["#FFF", "#12345", "#GGGGGG", "red", "#1234567"])
def test_unsupported(tmp_path, color):
    path = workbook(tmp_path / "bad.xlsx", color)
    before = path.read_bytes(), path.stat().st_mtime_ns
    with pytest.raises(ExtractionError) as error:
        _xlsx(path, tmp_path / "cache")
    assert error.value.code == "xlsx_unsupported_color"
    assert error.value.__cause__ is not None
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before
    assert not list(tmp_path.glob("cache/xlsx-*"))


def test_retry_cleanup_and_byte_preservation(tmp_path):
    path = workbook(tmp_path / "bad.xlsx", "#9EB3C9")
    original = path.read_bytes()
    cache = tmp_path / "cache"
    real_load = load_workbook

    def fail_retry(candidate, **kwargs):
        if candidate != path:
            assert candidate.parent.stat().st_mode & 0o777 == 0o700
            with ZipFile(path) as source, ZipFile(candidate) as repaired:
                assert source.namelist() == repaired.namelist()
                for info in source.infolist():
                    assert info.compress_type == repaired.getinfo(info.filename).compress_type
                    data = source.read(info)
                    expected = (
                        data.replace(b'rgb="#9EB3C9"', b'rgb="9EB3C9"')
                        if info.filename == compat.STYLES_PART
                        else data
                    )
                    assert repaired.read(info.filename) == expected
            raise ValueError("private path and contents")
        return real_load(candidate, **kwargs)

    with (
        patch("openpyxl.load_workbook", side_effect=fail_retry),
        pytest.raises(ExtractionError) as e,
    ):
        _xlsx(path, cache)
    assert e.value.code == "xlsx_stylesheet_unreadable"
    assert "private" not in str(e.value)
    assert not list(cache.iterdir()) and path.read_bytes() == original


@pytest.mark.parametrize("kind", ["corrupt", "worksheet", "missing"])
def test_unrelated(tmp_path, kind):
    path = workbook(tmp_path / "bad.xlsx")
    if kind == "corrupt":
        path.write_bytes(b"not a zip")
    else:
        rewrite(
            path,
            lambda n, d: (
                (b"<broken" if kind == "worksheet" else None)
                if n == "xl/worksheets/sheet1.xml"
                else d
            ),
        )
    with (
        patch("local_rag.extract.repair_styles", side_effect=AssertionError("no repair")),
        pytest.raises(ExtractionError) as e,
    ):
        _xlsx(path, tmp_path / "cache")
    assert e.value.code == "xlsx_parse_failed"


@pytest.mark.parametrize("declaration", [b"<!DOCTYPE styleSheet>", b'<!ENTITY x "secret">'])
def test_declarations(tmp_path, declaration):
    path = workbook(tmp_path / "bad.xlsx", "#9EB3C9")
    rewrite(path, lambda n, d: declaration + d if n == compat.STYLES_PART else d)
    with pytest.raises(ExtractionError) as e:
        _xlsx(path)
    assert e.value.code == "xlsx_stylesheet_unreadable"


@pytest.mark.parametrize("limit", ["members", "styles", "ratio", "total"])
def test_limits(tmp_path, limit, monkeypatch):
    path = workbook(tmp_path / "bad.xlsx", "#9EB3C9")
    if limit == "members":
        with ZipFile(path, "a") as archive:
            for i in range(10_001):
                archive.writestr(f"empty/{i}", b"")
    elif limit == "styles":
        with ZipFile(path) as archive:
            data = archive.read(compat.STYLES_PART)
        data += b" " * (compat.MAX_STYLES_BYTES + 1)
        rewrite(path, lambda n, d: data if n == compat.STYLES_PART else d)
        # Ensure this tests the style size rather than the compression ratio.
        with ZipFile(path) as archive:
            members = [(i.filename, archive.read(i)) for i in archive.infolist()]
        with ZipFile(path, "w", compression=ZIP_STORED) as archive:
            for name, data in members:
                archive.writestr(name, data)
    elif limit == "ratio":
        with ZipFile(path, "a", compression=ZIP_DEFLATED) as archive:
            archive.writestr("unrelated", b"a" * 1_000_000)
    else:
        monkeypatch.setattr(compat, "MAX_FILE_BYTES", 100)
    with pytest.raises(ExtractionError) as e:
        _xlsx(path)
    assert e.value.code == "xlsx_archive_limit"


def test_namespace_rewrite_only():
    data = (
        b'<s:styleSheet xmlns:s="http://schemas.openxmlformats.org/'
        b'spreadsheetml/2006/main" xmlns:x="urn:other"><!-- rgb="#abcdef" -->'
        b"<s:color rgb='#9eB3c9'/><x:color rgb=\"#bad\"/></s:styleSheet>"
    )
    repaired, count = compat.repair_styles(data)
    assert count == 1
    assert repaired == data.replace(b"rgb='#9eB3c9'", b"rgb='9eB3c9'")
    bad = data.replace(b"</s:styleSheet>", b'<s:color rgb="#FFF"/></s:styleSheet>')
    with pytest.raises(ExtractionError) as e:
        compat.repair_styles(bad)
    assert e.value.code == "xlsx_unsupported_color"


def test_drive_and_local_reports_and_target_retry(tmp_path):
    bad = workbook(tmp_path / "bad.xlsx", "#FFF")
    fixed = workbook(tmp_path / "repair.xlsx", "#9EB3C9")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    backend = TreeDrive({"root": [value("synthetic-file", "statement.xlsx", mime=mime)]})
    backend.content["synthetic-file"] = bad.read_bytes()
    home = tmp_path / "home"
    service = MultiSourceRAG(
        Settings(root=home, home=home), drive_backend_factory=lambda source: backend
    )
    service.add_drive_source("fixture", "root", "fixture-account", tmp_path / "token.json")
    report = service.reconcile()["sources"][0]
    assert report["errors"] == [
        "file 'synthetic-file': stage=parse error=ExtractionError code=xlsx_unsupported_color"
    ]
    coverage = service.index_coverage()
    assert coverage["issues"][0]["code"] == "xlsx_unsupported_color"
    reader = MCPServer(service, "reader").call("index_coverage", {})
    assert reader["issues"][0]["code"] is None and reader["issues"][0]["reason"] is None
    assert reader["issues"][0]["title"] is None
    backend.content["synthetic-file"] = fixed.read_bytes()
    backend.downloads.clear()
    report = service.reconcile("fixture", target="id:synthetic-file", force_index=True)["sources"][
        0
    ]
    assert not report["errors"] and report["indexed"] == 1
    assert backend.downloads == ["synthetic-file"]
    assert (
        report["warnings"][0] == "file 'synthetic-file': compatibility_repair=xlsx_rgb_hash_prefix"
    )
    document = service.db.document_snapshot(service.registry.get("fixture").id)[0]
    assert json.loads(document["metadata_json"])["compatibility_repairs"][0]["count"] == 1
    raw = list((home / "cache/sources").glob("*/raw/*.xlsx"))
    assert len(raw) == 1 and raw[0].read_bytes() == fixed.read_bytes()
    root = tmp_path / "local"
    root.mkdir()
    (root / "bad.xlsx").write_bytes(bad.read_bytes())
    service.add_local_source("local", root)
    report = service.reconcile("local")["sources"][0]
    assert report["errors"] == [
        "file 'bad.xlsx': stage=parse error=ExtractionError code=xlsx_unsupported_color"
    ]


@pytest.mark.parametrize(
    "payload",
    [
        b"<broken",
        b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<fills><fill><patternFill patternType="invalid"/></fill></fills></styleSheet>',
    ],
)
def test_other_stylesheet_errors(tmp_path, payload):
    path = workbook(tmp_path / "bad.xlsx")
    rewrite(path, lambda n, d: payload if n == compat.STYLES_PART else d)
    with (
        patch("local_rag.extract.repair_styles", side_effect=AssertionError("no repair")),
        pytest.raises(ExtractionError) as e,
    ):
        _xlsx(path)
    assert e.value.code == "xlsx_stylesheet_unreadable"


def test_normal_sync_retry_and_embedding_diagnostics(tmp_path):
    path = workbook(tmp_path / "bad.xlsx", "#FFF")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    backend = TreeDrive({"root": [value("retry-file", "retry.xlsx", mime=mime)]})
    backend.content["retry-file"] = path.read_bytes()

    class OfflineEmbeddings:
        identity = ("synthetic", "test")

        def embed(self, texts):
            raise RuntimeError("private provider error contents")

    home = tmp_path / "home"
    service = MultiSourceRAG(
        Settings(root=home, home=home),
        OfflineEmbeddings(),
        drive_backend_factory=lambda source: backend,
    )
    service.add_drive_source("fixture", "root", "fixture-account", tmp_path / "token.json")
    assert service.reconcile()["sources"][0]["errors"]
    assert service.registry.get("fixture").cursor is None
    backend.content["retry-file"] = workbook(path, "#9EB3C9").read_bytes()
    backend.downloads.clear()
    report = service.reconcile()["sources"][0]
    assert report["indexed"] == 1 and report["errors"] == []
    assert backend.downloads == ["retry-file"]
    assert any(w.startswith("embedding:") for w in report["warnings"])
    assert "private provider" not in json.dumps(report)
    metadata = service.metadata("drive:retry-file", source="fixture")
    assert metadata["automatic"]["compatibility_repairs"][0]["count"] == 1
    hits = service.search("Synthetic", source="fixture")["results"]
    assert hits[0]["metadata"]["compatibility_repairs"][0]["count"] == 1


@pytest.mark.parametrize("stage", ["download_export", "index"])
def test_drive_stage_separates_download_and_index(tmp_path, stage):
    from local_rag.drive import DriveItemError
    from local_rag.indexer import Indexer

    path = workbook(tmp_path / "valid.xlsx")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    backend = TreeDrive({"root": [value("synthetic", "book.xlsx", mime=mime)]})
    backend.content["synthetic"] = path.read_bytes()
    home = tmp_path / "home"
    service = MultiSourceRAG(
        Settings(root=home, home=home), drive_backend_factory=lambda source: backend
    )
    service.add_drive_source("fixture", "root", "fixture-account", tmp_path / "token.json")
    error = (
        DriveItemError("private download error")
        if stage == "download_export"
        else ValueError("private index error")
    )
    target = (
        patch.object(backend, "download", side_effect=error)
        if stage == "download_export"
        else patch.object(Indexer, "store_external", side_effect=error)
    )
    with target:
        report = service.reconcile()["sources"][0]
    assert report["errors"] == [f"file 'synthetic': stage={stage} error={type(error).__name__}"]
    assert "private" not in json.dumps(report)
    issue = service.index_coverage()["issues"][0]
    assert issue["stage"] == stage and issue["code"] is None
