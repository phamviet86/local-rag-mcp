from pathlib import Path
from unittest.mock import patch

from local_rag.config import Settings
from local_rag.service import MultiSourceRAG


def test_incomplete_local_discovery_never_deletes(tmp_path: Path):
    root = tmp_path / "documents"
    root.mkdir()
    nested = root / "nested"
    nested.mkdir()
    (nested / "note.txt").write_text("orchard harvest")
    home = tmp_path / "data"
    service = MultiSourceRAG(Settings(root=home, home=home))
    service.add_local_source("notes", root)
    service.reconcile()
    before = service.db.document_snapshot(service.registry.get("notes").id)
    for error in (FileNotFoundError("secret filename"), PermissionError("secret filename")):
        with patch("local_rag.indexer.os.scandir", side_effect=error):
            report = service.reconcile()
        assert report["sources"][0]["removed"] == 0
        assert report["sources"][0]["errors"]
        assert "secret filename" not in service.registry.get("notes").last_error
        assert [dict(row) for row in service.db.document_snapshot(before[0]["source_id"])] == [
            dict(row) for row in before
        ]
    import os

    real_scan = os.scandir

    def scan(path):
        if Path(path) == nested:
            raise PermissionError("private child")
        return real_scan(path)

    with patch("local_rag.indexer.os.scandir", side_effect=scan):
        assert service.reconcile()["sources"][0]["removed"] == 0
    moved = tmp_path / "relocated"
    root.rename(moved)
    assert service.reconcile()["sources"][0]["errors"]
    assert service.status()["documents"] == 1
    moved.rename(root)
    assert not service.reconcile()["sources"][0]["errors"]
    assert service.registry.get("notes").last_error is None


def test_local_moves_preserve_identity_and_do_not_extract(tmp_path: Path):
    root, home = tmp_path / "root", tmp_path / "data"
    root.mkdir()
    folder = root / "old"
    folder.mkdir()
    for i in range(2):
        (folder / f"note{i}.txt").write_text(f"orchard {i}")
    service = MultiSourceRAG(Settings(root=home, home=home))
    source = service.add_local_source("notes", root)
    service.reconcile()
    original = service.db.document_snapshot(source["id"])
    with service.db.transaction() as connection:
        connection.execute(
            """INSERT INTO relationships
            (source_document_id,target_document_id,relation,evidence_json,actor)
            VALUES(?,?,'related','[]','test')""",
            (original[0]["id"], original[1]["id"]),
        )
    folder.rename(root / "new")
    with patch.object(service.extractor, "extract", side_effect=AssertionError("no extraction")):
        report = service.reconcile()["sources"][0]
    assert report["moved"] == 2 and report["indexed"] == report["removed"] == 0
    updated = service.db.document_snapshot(source["id"])
    assert {row["id"] for row in updated} == {row["id"] for row in original}
    assert all(row["relative_path"].startswith("new/") for row in updated)
    with service.db.connect() as connection:
        assert connection.execute("SELECT count(*) FROM relationships").fetchone()[0] == 1
    # Editing during a move cannot transfer unrelated derived state.
    (root / "new/note0.txt").rename(root / "changed.txt")
    (root / "changed.txt").write_text("different content")
    report = service.reconcile()["sources"][0]
    assert report["moved"] == 0 and report["indexed"] == report["removed"] == 1
    (root / "copy.txt").write_text((root / "changed.txt").read_text())
    service.reconcile()
    (root / "copy.txt").rename(root / "copy-moved.txt")
    report = service.reconcile()["sources"][0]
    assert report["moved"] == 0 and report["indexed"] == report["removed"] == 1


def test_relocate_runs_as_admin_job_and_keeps_ids(tmp_path: Path):
    import pytest

    from local_rag.mcp import MCPServer

    root, home = tmp_path / "old", tmp_path / "home"
    root.mkdir()
    (root / "note.txt").write_text("orchard")
    service = MultiSourceRAG(Settings(root=home, home=home))
    source = service.add_local_source("notes", root)
    service.reconcile()
    old = service.db.document_snapshot(source["id"])[0]
    new = tmp_path / "new"
    root.rename(new)
    for profile in ("reader", "reviewer"):
        with pytest.raises(ValueError, match="not exposed"):
            MCPServer(service, profile).call(
                "relocate_source", {"source": "notes", "new_path": str(new)}
            )
    result = MCPServer(service, "admin").call(
        "relocate_source", {"source": "notes", "new_path": str(new)}
    )
    assert result["state"] == "complete"
    updated = service.db.document_snapshot(source["id"])[0]
    assert updated["id"] == old["id"] and updated["path"] == str(new / "note.txt")
    assert (new / "note.txt").read_text() == "orchard"
    result = service.start_index_job("relocate", "notes", str(root), background=False)
    assert result["state"] == "error"
    assert service.registry.get("notes").locator == str(new)


def test_stable_reference_survives_move_and_relocate(tmp_path: Path):
    import pytest

    root, home = tmp_path / "root", tmp_path / "home"
    root.mkdir()
    (root / "note.txt").write_text("orchard identity")
    service = MultiSourceRAG(Settings(root=home, home=home))
    source = service.add_local_source("notes", root)
    service.reconcile()
    hit = service.search("orchard", mode="full_text")["results"][0]
    ref = hit["stable_ref"]
    assert ref.startswith(f"local:{source['id']}:")
    (root / "note.txt").rename(root / "moved.txt")
    service.reconcile()
    assert service.read(ref)["path"] == "moved.txt"
    assert service.metadata(ref)["stable_ref"] == ref
    with pytest.raises(ValueError):
        service.read(ref, source="another-source")
    new = tmp_path / "relocated"
    root.rename(new)
    assert (
        service.start_index_job("relocate", "notes", str(new), background=False)["state"]
        == "complete"
    )
    assert service.read(ref)["text"] == "orchard identity"


def test_vietnamese_fts_keeps_original_text_and_vectors(tmp_path: Path):
    root, home = tmp_path / "root", tmp_path / "home"
    root.mkdir()
    (root / "note.txt").write_text("hợp đồng đường ĐƯỜNG")
    service = MultiSourceRAG(Settings(root=home, home=home))
    service.add_local_source("notes", root)
    service.reconcile()
    for query in ("hop", "hợp", "duong", "đường", "ĐƯỜNG"):
        hit = service.search(query, mode="full_text")["results"][0]
        assert hit["text"] == "hợp đồng đường ĐƯỜNG"
    evidence = [{"path": "notes:note.txt", "quote": "hợp"}]
    service.add_metadata("notes:note.txt", "tag", "đặc biệt", evidence, "test")
    assert service.search("dac", mode="full_text")["results"]


def test_hit_offsets_and_read_around_match_effective_text(tmp_path: Path):
    from local_rag.mcp import MCPServer

    root, home = tmp_path / "root", tmp_path / "home"
    root.mkdir()
    (root / "note.txt").write_text("orchard harvest " * 100)
    service = MultiSourceRAG(Settings(root=home, home=home, chunk_chars=100, chunk_overlap=10))
    service.add_local_source("notes", root)
    service.reconcile()
    for hit in service.search("orchard", mode="full_text")["results"]:
        if hit["chunk_ordinal"] is None:
            continue
        text = service.read(hit["stable_ref"])["text"]
        assert text[hit["start_char"] : hit["end_char"]] == hit["text"]
        assert hit["indexed_at"] and hit["source_modified_at"]
        expanded = MCPServer(service).call(
            "read", {"path": hit["stable_ref"], "around": hit["start_char"], "window": 150}
        )
        assert expanded["text"] == text[expanded["start"] : expanded["end"]]
        assert expanded["start"] <= hit["start_char"] < expanded["end"]


def test_diversify_and_collapse_keep_scoped_locations(tmp_path: Path):
    root, home = tmp_path / "root", tmp_path / "home"
    root.mkdir()
    text = "\n".join(f"orchard harvest section {i} " * 8 for i in range(10))
    (root / "one.txt").write_text(text)
    (root / "copy.txt").write_text(text)
    service = MultiSourceRAG(Settings(root=home, home=home, chunk_chars=100, chunk_overlap=10))
    service.add_local_source("notes", root)
    service.reconcile()
    results = service.search("orchard", limit=100, mode="full_text", per_document=1)["results"]
    refs = [hit["stable_ref"] for hit in results]
    assert len(refs) == len(set(refs))
    hashes = [hit["chunk_hash"] for hit in results if hit["chunk_hash"]]
    assert len(hashes) == len(set(hashes))
    assert any(
        {location["path"] for location in hit.get("locations", [])} == {"one.txt", "copy.txt"}
        for hit in results
    )
    results = service.search("orchard", folder="one.txt", mode="full_text")["results"]
    assert all(
        location["path"] == "one.txt" for hit in results for location in hit.get("locations", [])
    )


def test_docx_headings_keep_legacy_text_and_order_is_opt_in(tmp_path: Path):
    from docx import Document

    from local_rag.extract import _docx
    from local_rag.indexer import _chunks

    document = Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("before table")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "table contents"
    document.add_heading("Detail", level=2)
    document.add_paragraph("after table")
    path = tmp_path / "sample.docx"
    document.save(path)
    legacy = _docx(path)
    assert legacy.text == "Policy\n\nbefore table\n\nDetail\n\nafter table\n\ntable contents"
    assert legacy.spans[-1].metadata["heading_path"] == ["Policy"]
    assert legacy.spans[-2].metadata["heading_path"] == ["Policy", "Detail"]
    ordered = _docx(path, True)
    assert ordered.text.index("table contents") < ordered.text.index("after table")
    assert _chunks(legacy, 1200, 150)[0].chunk_hash != _chunks(ordered, 1200, 150)[0].chunk_hash


def test_semantic_blob_matches_json_and_invalidates_across_writers(tmp_path: Path):
    import json

    import pytest

    from tests.test_optimizations import BatchEmbeddings

    pytest.importorskip("numpy")
    root, home = tmp_path / "root", tmp_path / "home"
    root.mkdir()
    (root / "one.txt").write_text("orchard orchard")
    (root / "two.txt").write_text("ocean")
    service = MultiSourceRAG(Settings(root=home, home=home), BatchEmbeddings())
    source = service.add_local_source("notes", root)
    service.reconcile()
    engine = service.search_engine
    with patch("local_rag.search._numpy", return_value=None):
        baseline = engine._semantic("orchard", None, None, 100)
    accelerated = engine._semantic("orchard", None, None, 100)
    assert list(accelerated) == list(baseline)
    for key in baseline:
        assert accelerated[key][1] == pytest.approx(baseline[key][1], abs=1e-6)
    assert engine._matrix_cache
    (root / "new.txt").write_text("orchard ocean")
    other = MultiSourceRAG(service.settings, BatchEmbeddings())
    other.reconcile()
    assert len(engine._semantic("orchard", None, None, 100)) == 3
    # Another connection updating JSON refreshes BLOB and invalidates the cached matrix.
    with other.db.transaction() as connection:
        connection.execute("UPDATE vectors SET vector_json=?", (json.dumps([0, 0, 1]),))
    with patch("local_rag.search._numpy", return_value=None):
        baseline = engine._semantic("orchard", None, None, 100)
    accelerated = engine._semantic("orchard", None, None, 100)
    assert [value[1] for value in accelerated.values()] == pytest.approx(
        [value[1] for value in baseline.values()], abs=1e-6
    )
    other.enable_source(source["id"], False)
    with pytest.raises(RuntimeError, match="no cached vectors"):
        engine._semantic("orchard", None, None, 100)


def test_search_filters_apply_through_sdk_and_cli(tmp_path: Path, capsys):
    import json
    import os

    import anyio
    import pytest

    from local_rag.cli import main
    from local_rag.mcp import create_sdk_server

    root, home = tmp_path / "root", tmp_path / "home"
    root.mkdir()
    for name, timestamp in (
        ("old.txt", 1577836800),
        ("new.txt", 1735689600),
        ("new.md", 1735689600),
    ):
        path = root / name
        path.write_text(f"orchard {name}")
        os.utime(path, (timestamp, timestamp))
    settings = Settings(root=home, home=home)
    settings.save()
    service = MultiSourceRAG(settings)
    service.add_local_source("notes", root)
    service.reconcile()
    filters = {
        "media_type": "text/plain",
        "modified_after": "2024-01-01",
        "modified_before": "2026-01-01",
        "per_document": 1,
    }
    result = service.search("orchard", mode="full_text", **filters)
    assert [hit["path"] for hit in result["results"]] == ["new.txt"]

    async def check():
        server = create_sdk_server(service)
        result = await server.call_tool(
            "search", {"query": "orchard", "mode": "full_text", **filters}
        )
        assert not result.is_error
        assert [hit["path"] for hit in result.structured_content["results"]] == ["new.txt"]

    anyio.run(check)
    main(
        [
            "--home",
            str(home),
            "search",
            "orchard",
            "--mode",
            "full_text",
            "--media-type",
            "text/plain",
            "--modified-after",
            "2024-01-01",
            "--modified-before",
            "2026-01-01",
            "--per-document",
            "1",
        ]
    )
    assert [hit["path"] for hit in json.loads(capsys.readouterr().out)["results"]] == ["new.txt"]
    with pytest.raises(ValueError, match="must precede"):
        service.search("orchard", modified_after="2026-01-01", modified_before="2020-01-01")


def test_populated_v090_migrations_preserve_rows_vectors_and_chunk_text(tmp_path: Path):
    import json
    import sqlite3

    from local_rag.db import Database
    from local_rag.indexer import _chunks, _document_from_artifact

    fixture = Path(__file__).parent / "fixtures/v090-populated.json"
    payload = json.loads(fixture.read_text().replace("@FIXTURE@", str(tmp_path)))
    assert payload["generated_by_commit"] == "d9ec11b3f8cb2f222381b15583b6842070106939"
    home = tmp_path / "home"
    home.mkdir()
    db = Database(home / "index.sqlite3")
    with sqlite3.connect(db.path) as connection:
        connection.execute("CREATE TABLE schema_migrations(version INTEGER PRIMARY KEY)")
        for number, migration in enumerate(payload["migrations"], 1):
            connection.executescript(migration)
            connection.execute("INSERT INTO schema_migrations VALUES(?)", (number,))
        for table, rows in payload["tables"].items():
            for row in rows:
                columns = ",".join(row)
                placeholders = ",".join("?" for _ in row)
                connection.execute(
                    f"INSERT INTO {table}({columns}) VALUES({placeholders})", tuple(row.values())
                )
    for path, artifact in payload["artifacts"].items():
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(artifact))

    class NoEmbeddings:
        identity = ("fixture", "v1")

        def embed(self, texts):
            raise AssertionError("upgrade must not call embeddings")

    service = MultiSourceRAG(Settings(root=tmp_path / "root", home=home), NoEmbeddings())
    db.migrate()  # A second startup must be idempotent.
    with db.connect() as connection:
        for table in (
            "documents",
            "chunks",
            "vectors",
            "reviews",
            "review_revisions",
            "agent_metadata",
            "relationships",
        ):
            expected = payload["tables"][table]
            columns = ",".join(expected[0])
            actual = [dict(row) for row in connection.execute(f"SELECT {columns} FROM {table}")]
            assert actual == expected, table
        assert connection.execute(
            "SELECT count(*) FROM vectors WHERE vector_blob IS NOT NULL"
        ).fetchone()[0] == len(payload["tables"]["vectors"])
    assert service.read("default:scan.pdf")["has_review_corrections"]
    assert service.search("duong", mode="full_text")["results"]
    for document in payload["tables"]["documents"]:
        path = document["effective_artifact_path"] or document["artifact_path"]
        extracted = _document_from_artifact(payload["artifacts"][path])
        previous = [
            chunk for chunk in payload["tables"]["chunks"] if chunk["document_id"] == document["id"]
        ]
        assert [(chunk.text, chunk.chunk_hash) for chunk in _chunks(extracted, 1200, 150)] == [
            (chunk["text"], chunk["chunk_hash"]) for chunk in previous
        ]


def test_local_move_keeps_ocr_corrections_and_never_embeds(tmp_path: Path):
    from tests.helpers import write_text_pdf
    from tests.test_optimizations import BatchEmbeddings

    root, home = tmp_path / "root", tmp_path / "home"
    folder = root / "old"
    folder.mkdir(parents=True)
    write_text_pdf(folder / "scan.pdf", "Original uncertain page")
    embeddings = BatchEmbeddings()
    service = MultiSourceRAG(Settings(root=home, home=home, ocr_mode="no-ocr"), embeddings)
    service.add_local_source("notes", root)
    service.reconcile()
    review = service.reviews()[0]
    evidence = [{"path": "notes:old/scan.pdf", "locator": "page:1", "quote": "source"}]
    service.correct_review(review["id"], "Human corrected orchard", evidence, "reviewer")
    service.add_metadata("notes:old/scan.pdf", "tag", "kept", evidence, "reviewer")
    service.reconcile()
    before = service.db.resolve_document("notes:old/scan.pdf")
    calls = embeddings.calls
    folder.rename(root / "new")
    with patch.object(service.extractor, "extract", side_effect=AssertionError("must not extract")):
        report = service.reconcile()["sources"][0]
    assert report["moved"] == 1 and embeddings.calls == calls
    after = service.db.resolve_document("notes:new/scan.pdf")
    assert after["id"] == before["id"]
    assert after["effective_artifact_path"] == before["effective_artifact_path"]
    assert service.db.review(review["id"])["status"] == "resolved"
    assert service.metadata("notes:new/scan.pdf")["agent"][0]["key"] == "tag"
    hit = service.search("corrected", mode="full_text")["results"][0]
    assert (
        service.read(hit["stable_ref"])["text"][hit["start_char"] : hit["end_char"]] == hit["text"]
    )
