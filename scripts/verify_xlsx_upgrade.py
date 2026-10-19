#!/usr/bin/env python3
"""Verify two installed wheels against a synthetic 0.9.0 home through the CLI.

Run outside the checkout with --old-executable and --new-executable. A loopback-only
embedding stub records inputs; no production home, Drive account or external API is used.
"""

from __future__ import annotations

import argparse
import json
import os
import sqlite3
import subprocess
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import PatternFill


def verify(old: Path, new: Path, base: Path) -> dict[str, Any]:
    home, root = base / "home", base / "source"
    root.mkdir()
    for i in range(3):
        (root / f"note{i}.txt").write_text(f"Synthetic existing orchard note number {i}.")
    xlsx = root / "statement.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    sheet["A1"] = "FixtureHeader"
    sheet["A1"].fill = PatternFill("solid", fgColor="FF9EB3C9")
    sheet["B2"] = 123
    book.save(xlsx)
    book.close()
    with ZipFile(xlsx) as archive:
        members = [(info, archive.read(info)) for info in archive.infolist()]
    with ZipFile(xlsx, "w") as archive:
        for info, data in members:
            if info.filename == "xl/styles.xml":
                data = data.replace(b'rgb="FF9EB3C9"', b'rgb="#9EB3C9"')
            archive.writestr(info, data)
    original = xlsx.read_bytes(), xlsx.stat().st_mtime_ns
    inputs: list[str] = []

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, format: str, *args: Any) -> None:
            pass

        def do_POST(self) -> None:
            assert self.path == "/embeddings"
            payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            inputs.extend(payload["input"])
            response = json.dumps(
                {
                    "data": [
                        {"index": i, "embedding": [1.0, float(len(text)), 0.5]}
                        for i, text in enumerate(payload["input"])
                    ]
                }
            ).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(response)))
            self.end_headers()
            self.wfile.write(response)

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    env = {
        k: v
        for k, v in os.environ.items()
        if not k.startswith(("LOCAL_RAG_", "LOCAL_RAG_MCP_")) and k != "PYTHONPATH"
    }
    env.update(
        {
            "LOCAL_RAG_MCP_EMBEDDING_PROVIDER": "openai",
            "LOCAL_RAG_MCP_EMBEDDING_MODEL": "synthetic-upgrade",
            "LOCAL_RAG_MCP_OPENAI_API_KEY": "synthetic-loopback-only",
            "LOCAL_RAG_MCP_OPENAI_BASE_URL": f"http://127.0.0.1:{server.server_port}",
            "NO_PROXY": "127.0.0.1,localhost",
        }
    )

    def run(executable: Path, *args: str, allowed: tuple[int, ...] = (0,)) -> Any:
        result = subprocess.run(
            [str(executable), "--home", str(home), *args],
            capture_output=True,
            text=True,
            env=env,
            cwd=base,
        )
        assert result.returncode in allowed, (args, result.returncode, result.stderr, result.stdout)
        return json.loads(result.stdout)

    def snapshot() -> dict[str, Any]:
        with sqlite3.connect(home / "index.sqlite3") as connection:
            return {
                "documents": connection.execute(
                    "SELECT id,content_hash,artifact_path FROM documents ORDER BY id"
                ).fetchall(),
                "chunks": connection.execute(
                    "SELECT id,document_id,chunk_hash,text FROM chunks ORDER BY id"
                ).fetchall(),
                "vectors": connection.execute(
                    "SELECT chunk_hash,provider,model,vector_json FROM vectors ORDER BY chunk_hash"
                ).fetchall(),
                "migrations": connection.execute(
                    "SELECT version FROM schema_migrations ORDER BY version"
                ).fetchall(),
            }

    try:
        run(old, "setup", "--no-ocr")
        run(old, "source", "add-local", "fixture", str(root))
        run(old, "sync", "--source", "fixture", "--wait", allowed=(0, 1, 2))
        before = snapshot()
        assert len(before["documents"]) == len(before["chunks"]) == len(before["vectors"]) == 3
        assert before["migrations"][-1] == (4,)
        backup = base / "backup"
        subprocess.run(["cp", "-a", str(home), str(backup)], check=True)
        old_inputs = list(inputs)
        run(new, "doctor", "--json", allowed=(0, 2))
        status = run(new, "status", allowed=(0, 2))
        run(new, "source", "list")
        migrated = snapshot()
        assert migrated["migrations"][-3:] == [(5,), (6,), (7,)]
        for key in ("documents", "chunks", "vectors"):
            assert migrated[key] == before[key]
        assert inputs == old_inputs
        assert status["documents"] == 3
        report = run(new, "reindex", "--source", "fixture", "--target", "statement.xlsx", "--wait")
        assert report["state"] == "complete", report
        after = snapshot()
        assert len(after["documents"]) == 4
        for key in ("documents", "chunks", "vectors"):
            assert all(row in after[key] for row in before[key])
        new_inputs = inputs[len(old_inputs) :]
        assert new_inputs and all("FixtureHeader" in text for text in new_inputs)
        assert len(after["vectors"]) - len(before["vectors"]) == len(new_inputs)
        result = run(new, "search", "FixtureHeader", "--source", "fixture", "--mode", "full_text")
        hit = result["results"][0]
        metadata = run(new, "metadata", "get", hit["stable_ref"], "--source", "fixture")
        assert metadata["automatic"]["compatibility_repairs"][0]["count"] == 1
        assert "sheet:Synthetic/cell:A1" in json.dumps(hit)
        assert original == (xlsx.read_bytes(), xlsx.stat().st_mtime_ns)
        assert not list((home / "cache").glob("xlsx-*"))
        # Rollback restores the whole quiet data home, never downgrades schema in place.
        subprocess.run(["mv", str(home), str(base / "upgraded-home")], check=True)
        subprocess.run(["cp", "-a", str(backup), str(home)], check=True)
        run(old, "status", allowed=(0, 2))
        assert snapshot() == before
        return {
            "rollback_restored": True,
            "documents_before": 3,
            "documents_after": 4,
            "vectors_preserved": len(before["vectors"]),
            "new_vectors": len(new_inputs),
            "migrations": [5, 6, 7],
            "source_unchanged": True,
            "unrelated_documents_reembedded": 0,
        }
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--old-executable", required=True, type=Path)
    parser.add_argument("--new-executable", required=True, type=Path)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="rag-xlsx-upgrade-") as directory:
        print(
            json.dumps(
                verify(
                    args.old_executable.resolve(), args.new_executable.resolve(), Path(directory)
                ),
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
